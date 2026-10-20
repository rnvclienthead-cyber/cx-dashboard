#!/usr/bin/env python3
"""100 случайных отзывов: ответ конструктора и ИИ, как в разделе «Отзывы».

Не отправляет ничего на маркетплейсы. Только выгрузка для проверки менеджером.
Запуск на сервере:
  cd /root/cx-dashboard && PYTHONPATH=/root/cx-dashboard \\
    /root/cx-dashboard/venv/bin/python3 scripts/audit_review_auto_sample.py
"""
from __future__ import annotations

import os
import random
import sys
import time
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
os.chdir(ROOT)
sys.path.insert(0, str(ROOT))

from dotenv import load_dotenv

load_dotenv(ROOT / ".env")

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import text

from backend.database import SessionLocal
from backend.llm import complete_sync
from backend.qwen_quota import is_quota_or_model_error, resolve_free_qwen_models
from backend.routers import kb

SEED = 20260902
PER_BUCKET = 10  # на каждую оценку: 10 с текстом + 10 без
OUT_NAME = f"аудит_автоответов_отзывы_{datetime.now().strftime('%Y-%m-%d')}.xlsx"

SOURCE_SQL = """
SELECT mp, id, sku, rating, body, pros, cons, created_at, has_text FROM (
    SELECT
        'wb'::text AS mp,
        f.id::text AS id,
        COALESCE(f.supplier_article, '') AS sku,
        f.valuation AS rating,
        COALESCE(f.text, '') AS body,
        COALESCE(f.pros, '') AS pros,
        COALESCE(f.cons, '') AS cons,
        f.created_date AS created_at,
        (NULLIF(TRIM(COALESCE(f.text,'')), '') IS NOT NULL
         OR NULLIF(TRIM(COALESCE(f.pros,'')), '') IS NOT NULL
         OR NULLIF(TRIM(COALESCE(f.cons,'')), '') IS NOT NULL) AS has_text
    FROM wb_feedbacks f
    WHERE f.valuation BETWEEN 1 AND 5
    UNION ALL
    SELECT
        'ym',
        y.feedback_id::text,
        COALESCE(y.supplier_article, ''),
        y.valuation,
        COALESCE(y.comment, ''),
        COALESCE(y.pro_text, ''),
        COALESCE(y.contra_text, ''),
        y.created_date,
        (NULLIF(TRIM(COALESCE(y.comment,'')), '') IS NOT NULL
         OR NULLIF(TRIM(COALESCE(y.pro_text,'')), '') IS NOT NULL
         OR NULLIF(TRIM(COALESCE(y.contra_text,'')), '') IS NOT NULL)
    FROM ym_feedbacks y
    WHERE y.valuation BETWEEN 1 AND 5
    UNION ALL
    SELECT
        'ozon',
        o.review_id,
        COALESCE(o.supplier_article, ''),
        o.rating,
        COALESCE(o.review_text, ''),
        COALESCE(o.pros, ''),
        COALESCE(o.cons, ''),
        o.created_at,
        (NULLIF(TRIM(COALESCE(o.review_text,'')), '') IS NOT NULL
         OR NULLIF(TRIM(COALESCE(o.pros,'')), '') IS NOT NULL
         OR NULLIF(TRIM(COALESCE(o.cons,'')), '') IS NOT NULL)
    FROM ozon_reviews o
    WHERE o.rating BETWEEN 1 AND 5
) s
WHERE rating = :rating AND has_text = :has_text
ORDER BY random()
LIMIT :lim
"""


def _full_text(row: dict) -> str:
    parts = []
    for key in ("body", "pros", "cons"):
        val = (row.get(key) or "").strip()
        if val and val != "—":
            parts.append(val)
    return ". ".join(parts)


def _ai_draft(row: dict, reference: str) -> tuple[str, str]:
    """Тот же промпт, что /api/v1/cx-responses/draft для отзывов."""
    mp = row["mp"]
    mp_label = {"wb": "Wildberries", "ym": "Яндекс Маркет", "ozon": "Ozon"}.get(mp, mp)
    parts = []
    if row.get("sku"):
        parts.append(f"Артикул: {row['sku']}")
    if row.get("rating"):
        parts.append(f"Оценка: {row['rating']}/5")
    if (row.get("pros") or "").strip():
        parts.append(f"Плюсы: {row['pros']}")
    if (row.get("cons") or "").strip():
        parts.append(f"Минусы: {row['cons']}")
    if (row.get("body") or "").strip():
        parts.append(f"Текст: {row['body']}")
    context = "\n".join(parts) if parts else "Текст не указан"
    ref_block = ""
    if reference:
        ref_block = (
            "\n\nФирменный образец ответа — ориентируйся на его стиль, структуру "
            "и позицию бренда, но перефразируй по-своему, не копируй дословно:\n"
            f"«{reference}»"
        )
    system = (
        "Ты менеджер клиентского сервиса бренда товаров для дома Vidovito на маркетплейсе. "
        "Пиши живо, тепло и по-человечески, без канцелярита и клише. "
        "Отвечай только на русском языке. "
        "Структура ответа: сначала эмпатия/благодарность, затем решение или полезный совет, в конце — тёплое пожелание. "
        "Для негативных отзывов (1-3★) — искренне сожалей и предлагай конкретное решение (возврат по браку, «Чат с продавцом»). "
        "Для позитивных (4-5★) — благодари и добавляй лёгкий совет по эксплуатации. "
        "Не используй эмодзи. "
        "Не подписывайся и не добавляй прощание в конце — подпись добавится автоматически. "
        "Ответ должен быть готов к отправке — без заголовков и пояснений."
    )
    user_msg = (
        f"Маркетплейс: {mp_label}\n"
        f"Тип обращения: отзыв\n"
        f"{context}"
        f"{ref_block}\n\n"
        "Напиши ответ продавца."
    )
    models = ["qwen-plus"] + resolve_free_qwen_models("qwen-plus")
    last_err = ""
    for model in models:
        try:
            draft = complete_sync(
                system=system,
                user=user_msg,
                model_key=model,
                temperature=0.4,
                max_tokens=500,
            ).strip()
            draft += "\n\nАлександр, команда Видовито"
            return draft, model
        except Exception as e:
            last_err = str(e)
            if not is_quota_or_model_error(last_err):
                return f"[ошибка ИИ: {last_err}]", model
    return f"[ИИ недоступен: {last_err}]", "—"


def sample_reviews(db) -> list[dict]:
    rng = random.Random(SEED)
    picked: list[dict] = []
    used = set()
    for rating in (1, 2, 3, 4, 5):
        for has_text in (True, False):
            rows = db.execute(
                text(SOURCE_SQL),
                {"rating": rating, "has_text": has_text, "lim": 80},
            ).mappings().all()
            pool = []
            for r in rows:
                key = (r["mp"], r["id"])
                if key in used:
                    continue
                pool.append(dict(r))
            rng.shuffle(pool)
            take = pool[:PER_BUCKET]
            if len(take) < PER_BUCKET:
                extra = db.execute(
                    text(SOURCE_SQL.replace("AND has_text = :has_text", "")),
                    {"rating": rating, "has_text": has_text, "lim": 80},
                ).mappings().all()
                for r in extra:
                    key = (r["mp"], r["id"])
                    if key in used:
                        continue
                    if any(x["mp"] == r["mp"] and x["id"] == r["id"] for x in take):
                        continue
                    take.append(dict(r))
                    if len(take) >= PER_BUCKET:
                        break
            for r in take[:PER_BUCKET]:
                used.add((r["mp"], r["id"]))
                picked.append(r)
    rng.shuffle(picked)
    return picked


def write_xlsx(rows: list[dict], path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Отзывы"

    headers = [
        "№",
        "Площадка",
        "ID отзыва",
        "Артикул",
        "Оценка",
        "Есть текст",
        "Текст отзыва",
        "Достоинства",
        "Недостатки",
        "Дата",
        "Сценарий конструктора",
        "Нужна проверка",
        "Ответ конструктора (код)",
        "Ответ ИИ",
        "Модель ИИ",
        "Комментарий менеджера",
    ]
    header_font = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
    header_fill = PatternFill("solid", fgColor="BE123C")
    comment_fill = PatternFill("solid", fgColor="FEF3C7")
    wrap = Alignment(wrap_text=True, vertical="top")
    thin = Border(
        left=Side(style="thin", color="E5E7EB"),
        right=Side(style="thin", color="E5E7EB"),
        top=Side(style="thin", color="E5E7EB"),
        bottom=Side(style="thin", color="E5E7EB"),
    )

    for col, title in enumerate(headers, 1):
        cell = ws.cell(1, col, title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    for i, r in enumerate(rows, 1):
        values = [
            i,
            (r.get("mp") or "").upper(),
            r.get("id") or "",
            r.get("sku") or "",
            r.get("rating") or "",
            "да" if r.get("has_text") else "нет",
            r.get("body") or "",
            r.get("pros") or "",
            r.get("cons") or "",
            str(r.get("created_at") or "")[:19],
            r.get("scenario") or "",
            "да" if r.get("manual_check") else "нет",
            r.get("constructor") or "",
            r.get("ai") or "",
            r.get("ai_model") or "",
            "",
        ]
        for col, val in enumerate(values, 1):
            cell = ws.cell(i + 1, col, val)
            cell.alignment = wrap
            cell.border = thin
            if col == 16:
                cell.fill = comment_fill
        ws.row_dimensions[i + 1].height = 90

    widths = [5, 12, 22, 18, 10, 12, 36, 24, 24, 18, 28, 16, 42, 42, 16, 36]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 28

    info = wb.create_sheet("Как читать", 0)
    info["A1"] = "Как читать эту таблицу"
    info["A1"].font = Font(bold=True, size=14)
    lines = [
        "",
        "Это проверка автоответов на отзывы. На маркетплейсы ничего не уходило.",
        f"Выборка: {len(rows)} отзывов, seed={SEED}. По задумке — по 10 отзывов на каждую оценку 1–5, с текстом и без.",
        "",
        "Колонки:",
        "• Ответ конструктора — как кнопка «Из шаблона»: приветствие + эмпатия + решение + энергия + прощание.",
        "• Ответ ИИ — как кнопка «ИИ»: Qwen, опирается на тот же конструктор, но перефразирует.",
        "• Комментарий менеджера — жёлтое поле. Сюда напишите, что не так / как надо.",
        "",
        "После проверки верните этот файл как есть (xlsx). По комментариям доработаем шаблоны или промпт ИИ.",
        "",
        "Подсказки для комментария: «ок» / «нельзя так писать» / «не та проблема» / «слишком общее» / свой текст.",
    ]
    for i, line in enumerate(lines, 2):
        info.cell(i, 1, line)
        info.cell(i, 1).alignment = Alignment(wrap_text=True)
    info.column_dimensions["A"].width = 120

    wb.save(path)


def main() -> None:
    print("Выборка отзывов…", flush=True)
    db = SessionLocal()
    try:
        reviews = sample_reviews(db)
        print(f"Отобрано: {len(reviews)}", flush=True)
        for i, row in enumerate(reviews, 1):
            full = _full_text(row)
            signer = random.choice(kb.TEMPLATE_SIGNERS)
            answer, scenario, manual, _meta = kb.build_answer(
                db, row.get("sku") or "", full, row.get("rating") or 0, signer=signer,
            )
            row["constructor"] = answer or ""
            row["scenario"] = scenario or ""
            row["manual_check"] = bool(manual)
            if not answer:
                row["constructor"] = f"[конструктор не собрал ответ: {scenario}]"
            print(f"  [{i}/{len(reviews)}] {row['mp']} ★{row['rating']} конструктор → ИИ…", flush=True)
            ai_text, model = _ai_draft(row, answer or "")
            row["ai"] = ai_text
            row["ai_model"] = model
            time.sleep(0.15)
    finally:
        db.close()

    out = Path("/tmp") / OUT_NAME
    write_xlsx(reviews, out)
    print(f"Готово: {out}", flush=True)
    dist = {}
    for r in reviews:
        key = (r.get("rating"), "текст" if r.get("has_text") else "без текста", r.get("mp"))
        dist[key] = dist.get(key, 0) + 1
    print("Распределение (оценка, текст, мп):", flush=True)
    for k in sorted(dist):
        print(f"  {k}: {dist[k]}", flush=True)


if __name__ == "__main__":
    main()
