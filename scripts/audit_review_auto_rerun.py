#!/usr/bin/env python3
"""Пересобрать ответы конструктора и ИИ по уже проверенному файлу аудита.

Не отправляет ничего на маркетплейсы. Берёт те же отзывы и комментарии менеджера.
"""
from __future__ import annotations

import os
import sys
import time
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
os.chdir(ROOT)
sys.path.insert(0, str(ROOT))

from dotenv import load_dotenv

load_dotenv(ROOT / ".env")

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import text

from backend.database import SessionLocal
from backend.llm import complete_sync
from backend.qwen_quota import is_quota_or_model_error, resolve_free_qwen_models
from backend.review_policy import (
    MODE_AI_ONLY,
    MODE_APPROVE_TEMPLATE,
    MODE_AUTO_SORRY,
    MODE_AUTO_THANKS,
    REVIEW_AI_SYSTEM,
    apply_article_hints,
    article_facts_for_ai,
    classify_review_reply,
    no_text_answer,
    no_text_scenario,
    should_ground_constructor,
)
from backend.routers import kb

MODE_RU = {
    MODE_AUTO_THANKS: "авто: спасибо (4–5 без текста/фото)",
    MODE_AUTO_SORRY: "авто: извинение + чат (1–3 без текста/фото)",
    MODE_APPROVE_TEMPLATE: "черновик конструктора, отправка вручную",
    MODE_AI_ONLY: "пустое поле, ИИ только по кнопке",
}


def _full_text(row: dict) -> str:
    parts = []
    for key in ("body", "pros", "cons"):
        val = (row.get(key) or "").strip()
        if val and val != "—":
            parts.append(val)
    return ". ".join(parts)


def _ai_draft(row: dict, reference: str, mode: str) -> tuple[str, str]:
    mp = row["mp"]
    mp_label = {"wb": "Wildberries", "ym": "Яндекс Маркет", "ozon": "Ozon"}.get(mp, mp)
    parts = []
    if row.get("sku"):
        parts.append(f"Артикул: {row['sku']}")
        facts = article_facts_for_ai(row["sku"])
        if facts:
            parts.append(facts)
    if row.get("rating"):
        parts.append(f"Оценка: {row['rating']}/5")
    if (row.get("pros") or "").strip():
        parts.append(f"Плюсы: {row['pros']}")
    if (row.get("cons") or "").strip():
        parts.append(f"Минусы: {row['cons']}")
    if (row.get("body") or "").strip():
        parts.append(f"Текст: {row['body']}")
    context = "\n".join(parts) if parts else "Текст не указан"
    ref_block = ""
    if reference and should_ground_constructor(mode, row.get("body") or "", row.get("rating") or 0):
        ref_block = (
            f"\n\nФирменный образец ({row.get('scenario') or ''}) — возьми структуру и факты, "
            f"но перефразируй коротко, без лишних эмоций:\n«{reference}»"
        )
    user_msg = (
        f"Маркетплейс: {mp_label}\n"
        f"Тип обращения: отзыв\n"
        f"{context}"
        f"{ref_block}\n\n"
        "Напиши ответ продавца."
    )
    models = ["qwen-plus"] + resolve_free_qwen_models("qwen-plus")
    last_err = ""
    for model in models:
        try:
            draft = complete_sync(
                system=REVIEW_AI_SYSTEM,
                user=user_msg,
                model_key=model,
                temperature=0.4,
                max_tokens=500,
            ).strip()
            if row.get("sku"):
                draft = apply_article_hints(row["sku"], draft)
            draft += "\n\nАлександр, команда Видовито"
            return draft, model
        except Exception as e:
            last_err = str(e)
            if not is_quota_or_model_error(last_err):
                return f"[ошибка ИИ: {last_err}]", model
    return f"[ИИ недоступен: {last_err}]", "—"


def load_source(path: Path) -> list[dict]:
    wb = load_workbook(path, data_only=True)
    ws = wb["Отзывы"]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    rows = []
    for raw in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
        d = dict(zip(headers, raw))
        if not d.get("№"):
            continue
        rows.append({
            "n": d.get("№"),
            "mp": str(d.get("Площадка") or "").strip().lower(),
            "id": str(d.get("ID отзыва") or "").strip(),
            "sku": d.get("Артикул") or "",
            "rating": d.get("Оценка") or 0,
            "has_text": str(d.get("Есть текст") or "").strip().lower() in ("да", "true", "1"),
            "body": d.get("Текст отзыва") or "",
            "pros": d.get("Достоинства") or "",
            "cons": d.get("Недостатки") or "",
            "created_at": d.get("Дата") or "",
            "old_scenario": d.get("Сценарий конструктора") or "",
            "old_constructor": d.get("Ответ конструктора (код)") or "",
            "old_ai": d.get("Ответ ИИ") or "",
            "manager": d.get("Комментарий менеджера") or "",
        })
    return rows


def lookup_photos(db, mp: str, item_id: str) -> str:
    if mp == "wb":
        return db.execute(text(
            "SELECT COALESCE(photos,'') FROM wb_feedbacks WHERE id::text = :i"
        ), {"i": item_id}).scalar() or ""
    if mp == "ozon":
        return db.execute(text(
            "SELECT COALESCE(photos,'') FROM ozon_reviews WHERE review_id = :i"
        ), {"i": item_id}).scalar() or ""
    return ""


def write_xlsx(rows: list[dict], path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Отзывы"
    headers = [
        "№", "Площадка", "ID отзыва", "Артикул", "Оценка", "Есть текст",
        "Текст отзыва", "Достоинства", "Недостатки", "Дата",
        "Режим после правок", "Сценарий новый", "Нужна проверка",
        "Конструктор новый", "ИИ новый", "Модель ИИ",
        "Конструктор был", "ИИ был", "Комментарий менеджера",
    ]
    header_font = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
    header_fill = PatternFill("solid", fgColor="BE123C")
    new_fill = PatternFill("solid", fgColor="ECFDF5")
    comment_fill = PatternFill("solid", fgColor="FEF3C7")
    wrap = Alignment(wrap_text=True, vertical="top")
    thin = Border(
        left=Side(style="thin", color="E5E7EB"),
        right=Side(style="thin", color="E5E7EB"),
        top=Side(style="thin", color="E5E7EB"),
        bottom=Side(style="thin", color="E5E7EB"),
    )
    for col, title in enumerate(headers, 1):
        cell = ws.cell(1, col, title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    for i, r in enumerate(rows, 1):
        values = [
            r.get("n") or i,
            (r.get("mp") or "").upper(),
            r.get("id") or "",
            r.get("sku") or "",
            r.get("rating") or "",
            "да" if r.get("has_text") else "нет",
            r.get("body") or "",
            r.get("pros") or "",
            r.get("cons") or "",
            str(r.get("created_at") or "")[:19],
            r.get("mode_ru") or "",
            r.get("scenario") or "",
            "да" if r.get("manual_check") else "нет",
            r.get("constructor") or "",
            r.get("ai") or "",
            r.get("ai_model") or "",
            r.get("old_constructor") or "",
            r.get("old_ai") or "",
            r.get("manager") or "",
        ]
        for col, val in enumerate(values, 1):
            cell = ws.cell(i + 1, col, val)
            cell.alignment = wrap
            cell.border = thin
            if col in (14, 15):
                cell.fill = new_fill
            if col == 19:
                cell.fill = comment_fill
        ws.row_dimensions[i + 1].height = 96

    widths = [5, 10, 22, 18, 8, 10, 32, 20, 20, 16, 28, 24, 12, 42, 42, 14, 36, 36, 36]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 32

    info = wb.create_sheet("Как читать", 0)
    info["A1"] = "Повторная проверка после правок"
    info["A1"].font = Font(bold=True, size=14)
    lines = [
        "",
        "Те же 100 отзывов, что в файле с комментариями менеджера. На маркетплейсы ничего не уходило.",
        "Зелёные колонки — новые ответы. Жёлтая — прежний комментарий менеджера.",
        "",
        "Режимы:",
        "• авто: спасибо — 4–5 без текста и фото (само уйдёт, когда включите тумблер).",
        "• авто: извинение + чат — 1–3 без текста и фото, без возврата и замены сразу.",
        "• черновик конструктора — 4–5 с текстом или фото, отправка только вручную.",
        "• пустое поле, ИИ только по кнопке — 1–3 с текстом или фото. В очереди конструктор не подставляется.",
        "",
        "Сравните «новый» с комментарием менеджера. Если ок — можно выкладывать на сайт.",
    ]
    for i, line in enumerate(lines, 2):
        info.cell(i, 1, line)
        info.cell(i, 1).alignment = Alignment(wrap_text=True)
    info.column_dimensions["A"].width = 120
    wb.save(path)


def main() -> None:
    src = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("/tmp/аудит_менеджер.xlsx")
    print(f"Читаю {src}", flush=True)
    rows = load_source(src)
    print(f"Строк: {len(rows)}", flush=True)
    db = SessionLocal()
    try:
        for i, row in enumerate(rows, 1):
            photos = lookup_photos(db, row["mp"], row["id"])
            mode = classify_review_reply(
                row.get("rating"), row.get("body") or "", photos,
                row.get("pros") or "", row.get("cons") or "",
            )
            row["mode"] = mode
            row["mode_ru"] = MODE_RU.get(mode, mode)
            full = _full_text(row)
            if mode in (MODE_AUTO_THANKS, MODE_AUTO_SORRY):
                row["constructor"] = no_text_answer(row.get("rating"))
                row["scenario"] = no_text_scenario(row.get("rating"))
                row["manual_check"] = False
                reference = row["constructor"]
            elif mode == MODE_AI_ONLY:
                row["constructor"] = "— (в очереди поле пустое, только кнопка ИИ)"
                row["scenario"] = "только ИИ"
                row["manual_check"] = True
                reference = ""
            else:
                answer, scenario, manual, _meta = kb.build_answer(
                    db, row.get("sku") or "", full, row.get("rating") or 0,
                )
                row["constructor"] = answer or f"[конструктор не собрал ответ: {scenario}]"
                row["scenario"] = scenario or ""
                row["manual_check"] = bool(manual) or bool(photos)
                reference = answer or ""
            print(f"  [{i}/{len(rows)}] {row['mp']} ★{row['rating']} {mode} → ИИ…", flush=True)
            ai_text, model = _ai_draft(row, reference, mode)
            row["ai"] = ai_text
            row["ai_model"] = model
            time.sleep(0.15)
    finally:
        db.close()

    out = Path("/tmp") / f"аудит_автоответов_повтор_{datetime.now().strftime('%Y-%m-%d')}.xlsx"
    write_xlsx(rows, out)
    print(f"Готово: {out}", flush=True)


if __name__ == "__main__":
    main()
