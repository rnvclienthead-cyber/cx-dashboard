#!/usr/bin/env python3
"""Выгрузка: все SKU по 3 МП, индекс проблем как на главной, топ-10 тем обращений.

Запуск на сервере:
  cd /root/cx-dashboard && PYTHONPATH=/root/cx-dashboard \\
    /root/cx-dashboard/venv/bin/python3 scripts/export_sku_problems_xlsx.py
"""
from __future__ import annotations

import os
import re
import sys
from collections import defaultdict
from datetime import datetime, timedelta
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
os.chdir(ROOT)
sys.path.insert(0, str(ROOT))

from dotenv import load_dotenv

load_dotenv(ROOT / ".env")

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import text

from backend.database import SessionLocal

TOP_N = 10
PERIOD_DAYS = 90
MAX_QUOTES = 3
QUOTE_MAX = 280
# Ниже этого не показываем: шум, а не точка вмешательства.
MIN_SHARE_PCT = 2.0
MIN_COUNT = 3
# Актуальные SKU из ассортиментной матрицы: в продаже и новинки.
ACTIVE_MATRIX = ("В матрице", "Новинка")
SRC_LABEL = {
    "reviews": "отзыв",
    "questions": "вопрос",
    "chats": "чат",
    "claims": "претензия",
}
SELLER_RE = re.compile(
    r"благодарим вас|приносим.{0,12}извин|оформите.{0,50}возврат|"
    r"команда видовито|мы не имеем отношения|в данной ситуации просим|"
    r"здравствуйте\.\s+(мы|благодар)|перед упаковкой продукции",
    re.I,
)

CLAIM_CATS = {
    1: "Некомплект: фурнитура",
    2: "Некомплект: несущие детали",
    3: "Состояние упаковки",
    4: "Производственный дефект",
    5: "Механические повреждения",
    6: "Инструкция и сборка",
    7: "Хлипкость / устойчивость",
    8: "Пересорт / ошибка склада",
    9: "Качество материалов",
    10: "Габариты и размер",
    11: "Несоответствие описанию",
    12: "Субъективное «не подошло»",
    13: "Следы использования / б/у",
}

# Теги ИИ из отзывов → те же имена, что в претензиях, плюс нюансы, которые
# в 13 категориях не выделены отдельно.
AI_TAG_MAP = {
    "конструкция: хлипкость и неустойчивость": "Хлипкость / устойчивость",
    "конструкция: тонкий металл / пластик": "Качество материалов",
    "конструкция: слабые узлы и соединения": "Слабые узлы и соединения",
    "сборка: несовпадение пазов/отверстий": "Сборка: несовпадение пазов/отверстий",
    "сборка: непонятная инструкция": "Сборка: непонятная инструкция",
    "сборка: тяжелый физический монтаж": "Сборка: тяжёлый монтаж",
    "эргономика: мало места / вместимость": "Мало места / вместимость",
    "эргономика: несоответствие размерам вещей": "Несоответствие размерам вещей",
    "эргономика: неудобная форма изделия": "Неудобная форма",
    "ожидания: ощущение дешевизны вживую": "Ощущение дешевизны",
    "ожидания: отличие цвета от фото": "Цвет не как на фото",
    "ожидания: упаковка повреждена": "Состояние упаковки",
    "ожидания: упаковка": "Состояние упаковки",
    "функционал: запрос нового элемента": "Запрос доработки / нового элемента",
    "функционал: нет стопоров на колесах": "Нет стопоров на колёсах",
    "функционал: нехватка бортов/защиты": "Нехватка бортов / защиты",
}

POSITIVE_RE = re.compile(
    r"лёгк|легк|прост|удобн|вместительн|устойчив|хорош|отличн|"
    r"нрав|качественн|прочн|стабильн",
    re.I,
)
NEGATIVE_KEEP_RE = re.compile(
    r"неудобн|неустойчив|непрочн|некачеств|мало места|хлипк|слаб|тонк|"
    r"непонят|несовпад|тяж[её]л|дешев|отличие|поврежд|нет стопор|нехватка",
    re.I,
)

# (каноническое имя, ключевые куски текста)
KEYWORD_RULES: list[tuple[str, tuple[str, ...]]] = [
    ("Некомплект: фурнитура", (
        "некомплект", "недокомплект", "не хвата", "не долож", "нет болт",
        "нет гаек", "нет шуруп", "фурнитур", "крепеж", "нет заглуш",
    )),
    ("Некомплект: несущие детали", (
        "нет труб", "нет полк", "нет стоек", "нет ножк", "нет колес",
        "нет колёс", "несущ",
    )),
    ("Сборка: несовпадение пазов/отверстий", (
        "не совпад", "несовпад", "пазы", "пазов", "отверст", "не входит болт",
        "дырки не",
    )),
    ("Сборка: непонятная инструкция", (
        "инструкц", "непонятн", "схем сбор",
    )),
    ("Сборка: тяжёлый монтаж", (
        "трудно собр", "тяжело собр", "тяжёл", "тяжел собр", "молотком",
    )),
    ("Инструкция и сборка", (
        "сборк", "собрать", "собирает", "монтаж", "качает", "шатает",
    )),
    ("Хлипкость / устойчивость", (
        "хлипк", "хлюп", "неустойчив", "шатает", "качаетс", "развали",
        "провиса", "гнул", "гнёт",
    )),
    ("Слабые узлы и соединения", (
        "соединен", "узел", "сварк", "разболтал", "выпал",
    )),
    ("Качество материалов", (
        "тонкий металл", "тонкий пластик", "запах", "ржав", "краск",
        "облез", "материал", "дешев",
    )),
    ("Производственный дефект", (
        "брак", "дефект", "заводск", "крив", "скол", "трещин",
    )),
    ("Механические повреждения", (
        "вмятин", "погнут", "сломан", "сломал", "разбит", "царапин",
        "поврежд",
    )),
    ("Состояние упаковки", (
        "упаковк", "коробк", "помят", "разорв",
    )),
    ("Габариты и размер", (
        "размер", "габарит", "высота", "ширина", "глубина", "длина",
        "маленьк", "больш", "не влез", "не помещ",
    )),
    ("Мало места / вместимость", (
        "мало места", "вместим", "не влеза", "тесно",
    )),
    ("Несоответствие описанию", (
        "не соответствует", "не как на фото", "не как в описан",
        "обману", "другое описан",
    )),
    ("Цвет не как на фото", (
        "цвет", "оттен", "не тот цвет",
    )),
    ("Пересорт / ошибка склада", (
        "не тот товар", "пересорт", "прислали друго", "ошибка склада",
    )),
    ("Субъективное «не подошло»", (
        "не подош", "не оправд", "ожидан",
    )),
    ("Следы использования / б/у", (
        "б/у", "б у ", "следы использ", "уже пользова",
    )),
    ("Доставка / сроки", (
        "доставк", "курьер", "опоздал", "не пришел", "не приехал", "срок",
    )),
    ("Возврат / обмен", (
        "возврат", "вернуть", "обмен", "замен",
    )),
    ("Нет стопоров на колёсах", (
        "стопор", "не тормозит", "колёс", "колес",
    )),
]

_WB_NOT_PLAIN_REFUSAL = """(
    LOWER(TRIM(COALESCE(status::text, ''))) != 'отказ'
    OR COALESCE(status_ex, '') ILIKE '%доотправк%'
    OR EXISTS (
        SELECT 1 FROM reshipment_requests r
        WHERE r.matched_srid IS NOT NULL AND TRIM(r.matched_srid) <> ''
          AND LOWER(TRIM(r.matched_srid)) = LOWER(TRIM(wb_claims.srid))
    )
)"""
_CAT_ANY = (
    "(COALESCE(cat_1,false) OR COALESCE(cat_2,false) OR COALESCE(cat_3,false) "
    "OR COALESCE(cat_4,false) OR COALESCE(cat_5,false) OR COALESCE(cat_6,false) "
    "OR COALESCE(cat_7,false) OR COALESCE(cat_8,false) OR COALESCE(cat_9,false) "
    "OR COALESCE(cat_10,false) OR COALESCE(cat_11,false) OR COALESCE(cat_12,false) "
    "OR COALESCE(cat_13,false))"
)


def _in_period(ts, start, end) -> bool:
    if ts is None:
        return True
    try:
        if getattr(ts, "tzinfo", None) is not None:
            ts = ts.replace(tzinfo=None)
        return start <= ts <= end
    except Exception:
        return True


def _clean_sku(v) -> str:
    s = (str(v) if v is not None else "").strip()
    if not s or s.lower() in {"none", "без артикула", "nan"}:
        return ""
    return s


def _is_positive_tag(tag: str) -> bool:
    if NEGATIVE_KEEP_RE.search(tag):
        return False
    return bool(POSITIVE_RE.search(tag))


def canonicalize_ai_tag(tag: str) -> str | None:
    t = (tag or "").strip()
    if not t:
        return None
    if _is_positive_tag(t):
        return None
    mapped = AI_TAG_MAP.get(t.lower())
    if mapped:
        return mapped
    # Неизвестные теги оставляем как есть, без префикса КАТЕГОРИЯ:
    if ":" in t:
        _, rest = t.split(":", 1)
        rest = rest.strip()
        if rest and not _is_positive_tag(rest):
            return rest[0].upper() + rest[1:] if rest[0].islower() else rest
    return t


def classify_text(raw: str) -> str | None:
    t = (raw or "").lower()
    t = re.sub(r"\s+", " ", t).strip()
    if len(t) < 8:
        return None
    for name, kws in KEYWORD_RULES:
        if any(kw in t for kw in kws):
            return name
    return None


def _groups_from_ai(acc: SkuAcc, sku: str, raw_tags: list[str]) -> list[str]:
    groups = []
    seen = set()
    for t in raw_tags:
        g = canonicalize_ai_tag(t)
        if not g or g in seen:
            continue
        seen.add(g)
        groups.append(g)
        acc[sku].add_nuance(g, _tag_detail(t))
    return groups


def score_problem_sku(defects: int, orders: int, cost, rating) -> dict:
    rate = round(defects / orders * 1000, 1) if orders and orders > 0 else None
    if rate is None:
        f_norm = min(100.0, defects / 30 * 100)
    else:
        f_norm = min(100.0, rate / 40 * 100)
    v_norm = min(100.0, defects / 30 * 100)
    if cost is not None and cost > 0:
        money_norm = min(100.0, (defects * float(cost)) / 50000 * 100)
    else:
        money_norm = v_norm
    if rating is None:
        r_norm = 25.0
    elif rating >= 4.7:
        r_norm = 0.0
    else:
        r_norm = max(0.0, (4.7 - float(rating)) / 3.7 * 100)
    index = int(round(0.40 * f_norm + 0.25 * v_norm + 0.25 * money_norm + 0.10 * r_norm))
    index = max(0, min(100, index))
    if index < 30:
        zone = "низкий"
    elif index <= 60:
        zone = "средний"
    else:
        zone = "высокий"
    return {"rate_per_1000": rate, "problem_index": index, "problem_zone": zone}


def _table_exists(conn, name: str) -> bool:
    return bool(conn.execute(text(
        "SELECT 1 FROM information_schema.tables "
        "WHERE table_schema='public' AND table_name=:t"
    ), {"t": name}).scalar())


def _cols(conn, table: str) -> set[str]:
    rows = conn.execute(text(
        "SELECT column_name FROM information_schema.columns "
        "WHERE table_schema='public' AND table_name=:t"
    ), {"t": table}).fetchall()
    return {r[0] for r in rows}


def _prep_quote(raw: str, source: str) -> str | None:
    t = re.sub(r"\s+", " ", (raw or "")).strip()
    t = t.strip(" «»\"'")
    if len(t) < 25:
        return None
    if SELLER_RE.search(t):
        return None
    low = t.lower()
    if low in {"брак", "не подошло", "некачественный товар", "поврежден"}:
        return None
    if len(t) > QUOTE_MAX:
        cut = t[:QUOTE_MAX].rsplit(" ", 1)[0]
        t = (cut or t[:QUOTE_MAX]) + "…"
    label = SRC_LABEL.get(source, source)
    return f"[{label}] {t}"


def _tag_detail(tag: str) -> str:
    t = (tag or "").strip()
    if ":" in t:
        t = t.split(":", 1)[1].strip()
    return t


class SkuAcc:
    __slots__ = ("appeals", "problems", "src", "quotes", "nuances")

    def __init__(self):
        self.appeals = 0
        self.problems: dict[str, int] = defaultdict(int)
        self.src = {"reviews": 0, "questions": 0, "chats": 0, "claims": 0}
        self.quotes: dict[str, list[str]] = defaultdict(list)
        self.nuances: dict[str, dict[str, int]] = defaultdict(lambda: defaultdict(int))

    def add_nuance(self, name: str, detail: str | None):
        d = re.sub(r"\s+", " ", (detail or "")).strip(" .;,-")
        if not name or len(d) < 4:
            return
        if len(d) > 140:
            d = d[:139].rsplit(" ", 1)[0] + "…"
        if d.lower() == name.lower():
            return
        self.nuances[name][d[0].upper() + d[1:] if d[0].islower() else d] += 1

    def add_quote(self, name: str, raw: str | None, source: str):
        if not name or not raw:
            return
        q = _prep_quote(raw, source)
        if not q:
            return
        lst = self.quotes[name]
        if len(lst) >= MAX_QUOTES:
            return
        key = re.sub(r"\s+", " ", q.lower())[8:90]
        if any(key and key in x.lower() for x in lst):
            return
        lst.append(q)

    def add(self, source: str, names: list[str | None], quote: str | None = None):
        self.appeals += 1
        self.src[source] = self.src.get(source, 0) + 1
        seen = set()
        for n in names:
            if n and n not in seen:
                self.problems[n] += 1
                seen.add(n)
                self.add_quote(n, quote, source)

    def inside_text(self, name: str) -> str:
        items = sorted(self.nuances.get(name, {}).items(), key=lambda x: (-x[1], x[0]))
        if items:
            return "; ".join(d for d, _ in items[:6])
        quotes = self.quotes.get(name) or []
        if quotes:
            first = re.sub(r"^\[[^\]]+\]\s*", "", quotes[0])
            if len(first) > 180:
                first = first[:179].rsplit(" ", 1)[0] + "…"
            return first
        return ""


def _json_tags(raw) -> list[str]:
    if not raw:
        return []
    if isinstance(raw, str):
        import json
        try:
            raw = json.loads(raw)
        except Exception:
            return []
    if not isinstance(raw, dict):
        return []
    tags = raw.get("tags") or []
    return [str(t) for t in tags if t]


def load_platform(conn, mp: str, start, end) -> list[dict]:
    params = {"start": start, "end": end}
    acc: dict[str, SkuAcc] = defaultdict(SkuAcc)
    names: dict[str, str] = {}
    ratings: dict[str, float] = {}
    orders: dict[str, int] = {}
    costs: dict[str, float] = {}
    defects: dict[str, int] = {}
    sku_set: set[str] = set()
    nm_to_sku: dict[int, str] = {}
    active_matrix: set[str] = set()
    if _table_exists(conn, "wb_assortment"):
        for r in conn.execute(text("""
            SELECT TRIM(supplier_article) AS sku
            FROM wb_assortment
            WHERE matrix_status = ANY(:st)
              AND supplier_article IS NOT NULL
        """), {"st": list(ACTIVE_MATRIX)}).mappings():
            s = _clean_sku(r["sku"])
            if s:
                active_matrix.add(s)

    def remember(sku: str, name: str | None = None):
        sku = _clean_sku(sku)
        if not sku:
            return ""
        sku_set.add(sku)
        if name and sku not in names:
            names[sku] = str(name).strip()
        return sku

    # ── каталог / имена / рейтинг / заказы ────────────────────────────────
    if mp == "wb":
        if _table_exists(conn, "wb_ratings_totals"):
            for r in conn.execute(text("""
                SELECT TRIM(supplier_article) AS sku, average_rating, nm_id
                FROM wb_ratings_totals
                WHERE supplier_article IS NOT NULL
            """)).mappings():
                sku = remember(r["sku"])
                if sku and r["average_rating"] is not None:
                    ratings[sku] = round(float(r["average_rating"]), 2)
                try:
                    if sku and r["nm_id"] is not None:
                        nm_to_sku[int(r["nm_id"])] = sku
                except Exception:
                    pass
        if _table_exists(conn, "wb_assortment"):
            for r in conn.execute(text("""
                SELECT TRIM(supplier_article) AS sku, name_ru
                FROM wb_assortment WHERE supplier_article IS NOT NULL
            """)).mappings():
                remember(r["sku"], r["name_ru"])
        if _table_exists(conn, "wb_orders"):
            for r in conn.execute(text("""
                SELECT TRIM(supplier_article) AS sku, COUNT(*) AS cnt
                FROM wb_orders
                WHERE dt >= :start AND dt <= :end AND cancel_dt IS NULL
                  AND supplier_article IS NOT NULL AND TRIM(supplier_article) <> ''
                GROUP BY 1
            """), params).mappings():
                sku = remember(r["sku"])
                if sku:
                    orders[sku] = int(r["cnt"] or 0)

    elif mp == "ym":
        if _table_exists(conn, "ym_assortment"):
            for r in conn.execute(text("""
                SELECT TRIM(COALESCE(NULLIF(supplier_article,''), offer_id)) AS sku, name
                FROM ym_assortment
            """)).mappings():
                remember(r["sku"], r["name"])
        if _table_exists(conn, "ym_orders"):
            for r in conn.execute(text("""
                SELECT TRIM(supplier_article) AS sku,
                       MAX(sku_name) AS name,
                       COUNT(*) AS cnt
                FROM ym_orders
                WHERE created_at >= :start AND created_at <= :end
                  AND status NOT IN ('CANCELLED', 'RETURNED', 'PENDING_CANCELLED', 'PROCESSING_EXPIRED')
                  AND supplier_article IS NOT NULL AND TRIM(supplier_article) <> ''
                GROUP BY 1
            """), params).mappings():
                sku = remember(r["sku"], r["name"])
                if sku:
                    orders[sku] = int(r["cnt"] or 0)
        if _table_exists(conn, "ym_feedbacks"):
            for r in conn.execute(text("""
                SELECT TRIM(supplier_article) AS sku,
                       ROUND(AVG(valuation)::numeric, 2) AS rating
                FROM ym_feedbacks
                WHERE valuation IS NOT NULL AND supplier_article IS NOT NULL
                GROUP BY 1
            """)).mappings():
                sku = remember(r["sku"])
                if sku and r["rating"] is not None:
                    ratings[sku] = float(r["rating"])

    else:  # ozon
        if _table_exists(conn, "ozon_orders"):
            qty = "SUM(COALESCE(quantity,1))::bigint" if "quantity" in _cols(conn, "ozon_orders") else "COUNT(*)"
            for r in conn.execute(text(f"""
                SELECT TRIM(supplier_article) AS sku,
                       MAX(sku_name) AS name,
                       {qty} AS cnt
                FROM ozon_orders
                WHERE created_at >= :start AND created_at <= :end
                  AND supplier_article IS NOT NULL AND TRIM(supplier_article) <> ''
                  AND LOWER(COALESCE(status,'')) NOT IN ('cancelled','cancelled_from_split')
                GROUP BY 1
            """), params).mappings():
                sku = remember(r["sku"], r["name"])
                if sku:
                    orders[sku] = int(r["cnt"] or 0)
        if _table_exists(conn, "ozon_reviews"):
            for r in conn.execute(text("""
                SELECT TRIM(supplier_article) AS sku,
                       MAX(product_name) AS name,
                       ROUND(AVG(rating)::numeric, 2) AS rating
                FROM ozon_reviews
                WHERE supplier_article IS NOT NULL AND TRIM(supplier_article) <> ''
                  AND rating IS NOT NULL
                GROUP BY 1
            """)).mappings():
                sku = remember(r["sku"], r["name"])
                if sku and r["rating"] is not None:
                    ratings[sku] = float(r["rating"])

    try:
        for r in conn.execute(text("""
            SELECT DISTINCT ON (TRIM(supplier_article))
                TRIM(supplier_article) AS sku, cost_value
            FROM wb_cogs
            WHERE cost_value > 0
            ORDER BY TRIM(supplier_article), id DESC
        """)).mappings():
            sku = _clean_sku(r["sku"])
            if sku:
                costs[sku] = float(r["cost_value"] or 0)
    except Exception:
        pass

    # ── отзывы ────────────────────────────────────────────────────────────
    if mp == "wb" and _table_exists(conn, "wb_feedbacks"):
        rows = conn.execute(text("""
            SELECT TRIM(supplier_article) AS sku, valuation, text, pros, cons, ai_tags
            FROM wb_feedbacks
            WHERE created_date >= :start AND created_date <= :end
              AND supplier_article IS NOT NULL AND TRIM(supplier_article) <> ''
        """), params).mappings()
        for r in rows:
            sku = remember(r["sku"])
            if not sku:
                continue
            tags = _groups_from_ai(acc, sku, _json_tags(r["ai_tags"]))
            blob = " ".join(x for x in (r["cons"], r["text"], r["pros"]) if x)
            if not tags:
                rating = int(r["valuation"] or 0)
                if blob.strip() and rating <= 3:
                    tags = [classify_text(blob)]
                    if tags and tags[0]:
                        acc[sku].add_nuance(tags[0], blob)
            quote = (r["cons"] or "").strip() or (r["text"] or "").strip() or blob
            acc[sku].add("reviews", tags, quote=quote if tags else None)

    elif mp == "ym" and _table_exists(conn, "ym_feedbacks"):
        rows = conn.execute(text("""
            SELECT TRIM(supplier_article) AS sku, valuation, comment, pro_text, contra_text, ai_tags
            FROM ym_feedbacks
            WHERE created_date >= :start AND created_date <= :end
              AND supplier_article IS NOT NULL AND TRIM(supplier_article) <> ''
        """), params).mappings()
        for r in rows:
            sku = remember(r["sku"])
            if not sku:
                continue
            tags = _groups_from_ai(acc, sku, _json_tags(r["ai_tags"]))
            blob = " ".join(x for x in (r["contra_text"], r["comment"], r["pro_text"]) if x)
            if not tags:
                rating = int(r["valuation"] or 0)
                if blob.strip() and rating <= 3:
                    tags = [classify_text(blob)]
                    if tags and tags[0]:
                        acc[sku].add_nuance(tags[0], blob)
            quote = (r["contra_text"] or "").strip() or (r["comment"] or "").strip() or blob
            acc[sku].add("reviews", tags, quote=quote if tags else None)

    elif mp == "ozon" and _table_exists(conn, "ozon_reviews"):
        rows = conn.execute(text("""
            SELECT TRIM(supplier_article) AS sku, rating, review_text, pros, cons, product_name
            FROM ozon_reviews
            WHERE created_at >= :start AND created_at <= :end
              AND supplier_article IS NOT NULL AND TRIM(supplier_article) <> ''
        """), params).mappings()
        for r in rows:
            sku = remember(r["sku"], r["product_name"])
            if not sku:
                continue
            blob = " ".join(x for x in (r["cons"], r["review_text"], r["pros"]) if x)
            rating = int(r["rating"] or 0)
            tags = []
            if blob.strip() and (rating <= 3 or (r["cons"] or "").strip()):
                tags = [classify_text(blob)]
                if tags and tags[0]:
                    acc[sku].add_nuance(tags[0], blob)
            quote = (r["cons"] or "").strip() or (r["review_text"] or "").strip() or blob
            acc[sku].add("reviews", tags, quote=quote if tags else None)

    # ── вопросы ───────────────────────────────────────────────────────────
    if mp == "wb" and _table_exists(conn, "wb_questions"):
        for r in conn.execute(text("""
            SELECT TRIM(supplier_article) AS sku, text
            FROM wb_questions
            WHERE created_date >= :start AND created_date <= :end
              AND supplier_article IS NOT NULL AND TRIM(supplier_article) <> ''
              AND text IS NOT NULL AND TRIM(text) <> ''
        """), params).mappings():
            sku = remember(r["sku"])
            if sku:
                acc[sku].add("questions", [classify_text(r["text"])], quote=r["text"])

    elif mp == "ym" and _table_exists(conn, "ym_questions"):
        for r in conn.execute(text("""
            SELECT TRIM(supplier_article) AS sku, text
            FROM ym_questions
            WHERE created_date >= :start AND created_date <= :end
              AND supplier_article IS NOT NULL AND TRIM(supplier_article) <> ''
              AND text IS NOT NULL AND TRIM(text) <> ''
        """), params).mappings():
            sku = remember(r["sku"])
            if sku:
                acc[sku].add("questions", [classify_text(r["text"])], quote=r["text"])

    elif mp == "ozon" and _table_exists(conn, "ozon_questions"):
        date_col = "published_at" if "published_at" in _cols(conn, "ozon_questions") else "synced_at"
        for r in conn.execute(text(f"""
            SELECT TRIM(supplier_article) AS sku, text
            FROM ozon_questions
            WHERE {date_col} >= :start AND {date_col} <= :end
              AND supplier_article IS NOT NULL AND TRIM(supplier_article) <> ''
              AND text IS NOT NULL AND TRIM(text) <> ''
        """), params).mappings():
            sku = remember(r["sku"])
            if sku:
                acc[sku].add("questions", [classify_text(r["text"])], quote=r["text"])

    # ── чаты ──────────────────────────────────────────────────────────────
    if mp == "wb" and _table_exists(conn, "wb_chats"):
        for r in conn.execute(text("""
            SELECT nm_id, last_message, good_name, updated_at, created_at
            FROM wb_chats
            WHERE last_message IS NOT NULL AND TRIM(last_message) <> ''
        """)).mappings():
            ts = r["updated_at"] or r["created_at"]
            if not _in_period(ts, start, end):
                continue
            sku = ""
            try:
                sku = nm_to_sku.get(int(r["nm_id"] or 0), "")
            except Exception:
                sku = ""
            if not sku:
                continue
            remember(sku, r["good_name"])
            acc[sku].add("chats", [classify_text(r["last_message"])], quote=r["last_message"])

    elif mp == "ym" and _table_exists(conn, "ym_chats"):
        cols = _cols(conn, "ym_chats")
        art = "context_article" if "context_article" in cols else None
        if art:
            for r in conn.execute(text(f"""
                SELECT TRIM({art}) AS sku, last_message, updated_at, created_at
                FROM ym_chats
                WHERE last_message IS NOT NULL AND TRIM(last_message) <> ''
                  AND {art} IS NOT NULL AND TRIM({art}) <> ''
            """)).mappings():
                ts = r["updated_at"] or r["created_at"]
                if not _in_period(ts, start, end):
                    continue
                sku = remember(r["sku"])
                if sku:
                    acc[sku].add("chats", [classify_text(r["last_message"])], quote=r["last_message"])

    elif mp == "ozon" and _table_exists(conn, "ozon_chats"):
        for r in conn.execute(text("""
            SELECT TRIM(context_article) AS sku, last_message, updated_at, created_at
            FROM ozon_chats
            WHERE last_message IS NOT NULL AND TRIM(last_message) <> ''
              AND context_article IS NOT NULL AND TRIM(context_article) <> ''
        """)).mappings():
            ts = r["updated_at"] or r["created_at"]
            if not _in_period(ts, start, end):
                continue
            sku = remember(r["sku"])
            if sku:
                acc[sku].add("chats", [classify_text(r["last_message"])], quote=r["last_message"])

    # ── претензии / возвраты → индекс + темы ──────────────────────────────
    if mp == "wb" and _table_exists(conn, "wb_claims"):
        cat_sum = ",\n".join(
            f"SUM(CASE WHEN cat_{i} THEN 1 ELSE 0 END) AS c{i}" for i in range(1, 14)
        )
        rows = conn.execute(text(f"""
            SELECT TRIM(supplier_article) AS sku,
                   COUNT(*) FILTER (WHERE {_CAT_ANY}) AS defects,
                   {cat_sum}
            FROM wb_claims
            WHERE created_dt >= :start AND created_dt <= :end
              AND {_WB_NOT_PLAIN_REFUSAL}
              AND supplier_article IS NOT NULL
              AND TRIM(supplier_article) NOT IN ('', 'Без артикула')
            GROUP BY 1
        """), params).mappings()
        for r in rows:
            sku = remember(r["sku"])
            if not sku:
                continue
            defects[sku] = defects.get(sku, 0) + int(r["defects"] or 0)
            n = int(r["defects"] or 0)
            if n:
                acc[sku].src["claims"] += n
                acc[sku].appeals += n
                for i in range(1, 14):
                    cnt = int(r[f"c{i}"] or 0)
                    if cnt:
                        acc[sku].problems[CLAIM_CATS[i]] += cnt

    elif mp == "ym" and _table_exists(conn, "ym_returns"):
        cat_sum = ",\n".join(
            f"SUM(CASE WHEN cat_{i} THEN 1 ELSE 0 END) AS c{i}" for i in range(1, 14)
        )
        rows = conn.execute(text(f"""
            SELECT TRIM(supplier_article) AS sku,
                   COUNT(*) FILTER (WHERE {_CAT_ANY}) AS defects,
                   {cat_sum}
            FROM ym_returns
            WHERE created_at >= :start AND created_at <= :end
              AND supplier_article IS NOT NULL
              AND TRIM(supplier_article) NOT IN ('', 'Без артикула')
            GROUP BY 1
        """), params).mappings()
        for r in rows:
            sku = remember(r["sku"])
            if not sku:
                continue
            defects[sku] = defects.get(sku, 0) + int(r["defects"] or 0)
            n = int(r["defects"] or 0)
            if n:
                acc[sku].src["claims"] += n
                acc[sku].appeals += n
                for i in range(1, 14):
                    cnt = int(r[f"c{i}"] or 0)
                    if cnt:
                        acc[sku].problems[CLAIM_CATS[i]] += cnt

    elif mp == "ozon" and _table_exists(conn, "ozon_returns"):
        cat_sum = ",\n".join(
            f"SUM(CASE WHEN COALESCE(cat_{i},false) THEN 1 ELSE 0 END) AS c{i}"
            for i in range(1, 14)
        )
        rows = conn.execute(text(f"""
            SELECT TRIM(supplier_article) AS sku,
                   COUNT(*) FILTER (WHERE {_CAT_ANY}) AS defects,
                   {cat_sum}
            FROM ozon_returns
            WHERE COALESCE(return_date, synced_at) >= :start
              AND COALESCE(return_date, synced_at) <= :end
              AND supplier_article IS NOT NULL
              AND TRIM(supplier_article) NOT IN ('', 'Без артикула')
            GROUP BY 1
        """), params).mappings()
        for r in rows:
            sku = remember(r["sku"])
            if not sku:
                continue
            defects[sku] = defects.get(sku, 0) + int(r["defects"] or 0)
            n = int(r["defects"] or 0)
            if n:
                acc[sku].src["claims"] += n
                acc[sku].appeals += n
                for i in range(1, 14):
                    cnt = int(r[f"c{i}"] or 0)
                    if cnt:
                        acc[sku].problems[CLAIM_CATS[i]] += cnt

    cat_select = ",\n".join(f"COALESCE(cat_{i}, false) AS c{i}" for i in range(1, 14))
    if mp == "wb" and _table_exists(conn, "wb_claims"):
        for r in conn.execute(text(f"""
            SELECT TRIM(supplier_article) AS sku, user_comment AS comment, {cat_select}
            FROM wb_claims
            WHERE created_dt >= :start AND created_dt <= :end
              AND {_WB_NOT_PLAIN_REFUSAL}
              AND supplier_article IS NOT NULL
              AND TRIM(user_comment) <> ''
        """), params).mappings():
            sku = _clean_sku(r["sku"])
            if not sku:
                continue
            for i in range(1, 14):
                if r.get(f"c{i}"):
                    acc[sku].add_quote(CLAIM_CATS[i], r["comment"], "claims")
    elif mp == "ym" and _table_exists(conn, "ym_returns"):
        for r in conn.execute(text(f"""
            SELECT TRIM(supplier_article) AS sku, return_comment AS comment, {cat_select}
            FROM ym_returns
            WHERE created_at >= :start AND created_at <= :end
              AND supplier_article IS NOT NULL
              AND TRIM(COALESCE(return_comment,'')) <> ''
        """), params).mappings():
            sku = _clean_sku(r["sku"])
            if not sku:
                continue
            for i in range(1, 14):
                if r.get(f"c{i}"):
                    acc[sku].add_quote(CLAIM_CATS[i], r["comment"], "claims")
    elif mp == "ozon" and _table_exists(conn, "ozon_returns"):
        for r in conn.execute(text(f"""
            SELECT TRIM(supplier_article) AS sku, return_reason AS comment, {cat_select}
            FROM ozon_returns
            WHERE COALESCE(return_date, synced_at) >= :start
              AND COALESCE(return_date, synced_at) <= :end
              AND supplier_article IS NOT NULL
              AND TRIM(COALESCE(return_reason,'')) <> ''
        """), params).mappings():
            sku = _clean_sku(r["sku"])
            if not sku:
                continue
            for i in range(1, 14):
                if r.get(f"c{i}"):
                    acc[sku].add_quote(CLAIM_CATS[i], r["comment"], "claims")

    rows_out = []
    for sku in sku_set:
        if active_matrix and sku not in active_matrix:
            continue
        o = orders.get(sku, 0)
        if o <= 0:
            continue
        a = acc.get(sku) or SkuAcc()
        d = defects.get(sku, 0)
        metrics = score_problem_sku(d, o, costs.get(sku), ratings.get(sku))
        total = a.appeals
        ranked = sorted(a.problems.items(), key=lambda x: (-x[1], x[0]))
        top = []
        for name, cnt in ranked:
            share = round(cnt / total * 100, 1) if total else 0.0
            if share < MIN_SHARE_PCT or cnt < MIN_COUNT:
                continue
            cases = "\n".join(a.quotes.get(name) or [])
            top.append({
                "name": name,
                "count": cnt,
                "share": share,
                "inside": a.inside_text(name),
                "cases": cases,
            })
            if len(top) >= TOP_N:
                break
        if not top:
            continue
        rows_out.append({
            "sku": sku,
            "name": names.get(sku, ""),
            "rating": ratings.get(sku),
            "orders": o,
            "defects": d,
            "appeals": total,
            "src": a.src,
            "top": top,
            **metrics,
        })
    rows_out.sort(key=lambda x: (-x["problem_index"], -x["defects"], -x["appeals"], x["sku"]))
    return rows_out


def write_xlsx(by_mp: dict[str, list[dict]], start, end, path: Path) -> None:
    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
    header_fill = PatternFill("solid", fgColor="1F4E79")
    wrap = Alignment(wrap_text=True, vertical="center")
    thin = Border(
        left=Side(style="thin", color="D9E2F3"),
        right=Side(style="thin", color="D9E2F3"),
        top=Side(style="thin", color="D9E2F3"),
        bottom=Side(style="thin", color="D9E2F3"),
    )
    zone_fill = {
        "высокий": PatternFill("solid", fgColor="F4CCCC"),
        "средний": PatternFill("solid", fgColor="FFF2CC"),
        "низкий": PatternFill("solid", fgColor="D9EAD3"),
    }
    odd_fill = PatternFill("solid", fgColor="F7F9FC")

    legend = wb.active
    legend.title = "Как читать"
    legend["A1"] = "Проблемы SKU по площадкам"
    legend["A1"].font = Font(bold=True, size=16, color="1F4E79")
    lines = [
        "",
        f"Период: {start:%d.%m.%Y} — {end:%d.%m.%Y} (последние {PERIOD_DAYS} дней).",
        "Только актуальные артикулы: статус матрицы «В матрице» (в продаже) и «Новинка»,",
        "и были заказы на этой площадке за эти 3 месяца.",
        "Три листа: Wildberries, Яндекс Маркет, Ozon. Сортировка — по индексу проблем с главной (0–100).",
        "",
        "Индекс проблем — смесь: как часто брак на 1000 заказов, сколько случаев, деньги, рейтинг.",
        "Зона: низкий < 30, средний 30–60, высокий > 60.",
        "",
        "На каждом листе одна строка = одна группа у артикула. Повторяющийся артикул в колонке A склеен.",
        "В таблице только группы, которые реально тянут внимание:",
        f"доля ≥ {MIN_SHARE_PCT:g}% от всех обращений по артикулу и не меньше {MIN_COUNT} обращений.",
        "Мелочь на 1 упоминание из сотни сюда не попадает.",
        "Колонки: группа → доля → количество → что внутри → цитаты клиентов.",
        "",
        "«Что внутри» — конкретные нюансы (из тегов отзывов и формулировок клиентов), не только название группы.",
        "«Цитаты» — 2–3 живых обращения: отзыв, вопрос, претензия. Ответы бота не берём.",
        "Доля % = обращения с этой темой / все обращения по артикулу. Сумма долей может быть > 100%.",
        "",
        "Источники: отзывы (WB/ЯМ — теги ИИ, Ozon — по тексту), вопросы, чаты, претензии/возвраты.",
    ]
    for i, line in enumerate(lines, 2):
        legend[f"A{i}"] = line
        legend[f"A{i}"].alignment = Alignment(wrap_text=True)
        legend.row_dimensions[i].height = 18
    legend.column_dimensions["A"].width = 110

    sheets = [
        ("wb", "Wildberries"),
        ("ym", "Яндекс Маркет"),
        ("ozon", "Ozon"),
    ]
    headers = [
        "Артикул", "Название", "Индекс проблем", "Зона", "Рейтинг",
        "Заказы", "Всего обращений", "Место", "Группа", "Доля %", "Кол-во",
        "Что внутри группы", "Цитаты клиентов",
    ]

    for key, title in sheets:
        ws = wb.create_sheet(title)
        ws.sheet_properties.tabColor = {
            "wb": "7B3FF2", "ym": "FFCC00", "ozon": "005BFF",
        }[key]
        for col, h in enumerate(headers, 1):
            cell = ws.cell(1, col, h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.row_dimensions[1].height = 28
        ws.freeze_panes = "C2"

        data = by_mp.get(key) or []
        ri = 2
        sku_fill_a = PatternFill("solid", fgColor="E8F0FE")
        sku_fill_b = PatternFill("solid", fgColor="FFFFFF")
        for sku_i, row in enumerate(data):
            items = [x for x in (row["top"] or []) if x.get("name")]
            if not items:
                continue
            start_ri = ri
            block_fill = sku_fill_a if sku_i % 2 == 0 else sku_fill_b
            for rank, item in enumerate(items, 1):
                vals = [
                    row["sku"] if rank == 1 else "",
                    row["name"] if rank == 1 else "",
                    row["problem_index"] if rank == 1 else "",
                    row["problem_zone"] if rank == 1 else "",
                    (row["rating"] if row["rating"] is not None else "") if rank == 1 else "",
                    row["orders"] if rank == 1 else "",
                    row["appeals"] if rank == 1 else "",
                    rank,
                    item.get("name") or "",
                    item.get("share") if item.get("name") else "",
                    item.get("count") if item.get("name") else "",
                    item.get("inside") or "",
                    item.get("cases") or "",
                ]
                for col, val in enumerate(vals, 1):
                    cell = ws.cell(ri, col, val)
                    cell.alignment = wrap
                    cell.border = thin
                    cell.fill = block_fill
                    if col == 4 and rank == 1:
                        cell.fill = zone_fill.get(row["problem_zone"], block_fill)
                        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                ws.row_dimensions[ri].height = 72 if (item.get("cases") or item.get("inside")) else 20
                ri += 1
            end_ri = ri - 1
            if end_ri > start_ri:
                for col in range(1, 8):
                    letter = get_column_letter(col)
                    ws.merge_cells(f"{letter}{start_ri}:{letter}{end_ri}")
                    cell = ws.cell(start_ri, col)
                    cell.alignment = Alignment(
                        wrap_text=True,
                        vertical="center",
                        horizontal="center" if col >= 3 else "left",
                    )

        last = max(2, ri - 1)
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{last}"
        widths = [18, 34, 14, 12, 10, 12, 14, 8, 28, 10, 10, 48, 70]
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main() -> None:
    end = datetime.now().replace(hour=23, minute=59, second=59, microsecond=0)
    start = (end - timedelta(days=PERIOD_DAYS)).replace(hour=0, minute=0, second=0, microsecond=0)
    stamp = datetime.now().strftime("%Y-%m-%d")
    out = Path(os.environ.get("OUT_XLSX") or f"/tmp/SKU_проблемы_3МП_{stamp}.xlsx")

    print(f"Период {start:%Y-%m-%d} … {end:%Y-%m-%d}", flush=True)
    db = SessionLocal()
    by_mp = {}
    try:
        with db.bind.connect() as conn:
            for mp, label in (("wb", "WB"), ("ym", "ЯМ"), ("ozon", "Ozon")):
                print(f"Считаю {label}…", flush=True)
                rows = load_platform(conn, mp, start, end)
                by_mp[mp] = rows
                print(f"  {label}: {len(rows)} артикулов", flush=True)
    finally:
        db.close()

    write_xlsx(by_mp, start, end, out)
    print(f"Готово: {out}", flush=True)


if __name__ == "__main__":
    main()
